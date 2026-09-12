from __future__ import annotations

from pathlib import Path
import json

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _style_sheet(ws) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E1F2")
    border = Border(bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    ws.freeze_panes = "A2"

    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
            cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[column_letter].width = min(max_length + 2, 30)


def export_excel_workbooks(config: dict, outputs: dict, profile_tables: dict, confusion_tables: dict) -> None:
    tables_dir = Path(config["paths"]["tables_dir"])
    metrics_path = tables_dir / "metrics_summary.xlsx"
    rounds_path = tables_dir / "round_logs.xlsx"

    with pd.ExcelWriter(metrics_path, engine="openpyxl") as writer:
        for task_name, task_output in outputs["tasks"].items():
            task_output["comparison_table"].to_excel(writer, sheet_name=f"summary_{task_name}", index=False)
            pd.DataFrame(task_output["local_rows"]).drop(columns=["predictions", "truth", "confusion_matrix"], errors="ignore").to_excel(
                writer, sheet_name=f"local_{task_name}", index=False
            )
            pd.DataFrame(task_output["cross_device_rows"]).drop(columns=["predictions", "truth", "confusion_matrix"], errors="ignore").to_excel(
                writer, sheet_name=f"cross_{task_name}", index=False
            )
            confusion_tables[task_name].to_excel(writer, sheet_name=f"conf_{task_name}")
        for sheet_name, frame in profile_tables.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)

    with pd.ExcelWriter(rounds_path, engine="openpyxl") as writer:
        for task_name, task_output in outputs["tasks"].items():
            round_rows = []
            for round_row in task_output["federated_rounds"]:
                for client in round_row["clients"]:
                    round_rows.append({
                        "round": round_row["round"],
                        "task": round_row["task"],
                        "mode": round_row["mode"],
                        **round_row["metrics"],
                        **client,
                    })
            pd.DataFrame(round_rows).to_excel(writer, sheet_name=task_name[:31], index=False)

    for workbook_path in [metrics_path, rounds_path]:
        workbook = load_workbook(workbook_path)
        for sheet_name in workbook.sheetnames:
            _style_sheet(workbook[sheet_name])
        workbook.save(workbook_path)


def _safe_sheet_name(base: str, used: set[str]) -> str:
    candidate = base[:31]
    suffix = 1
    while candidate in used:
        tail = f"_{suffix}"
        candidate = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def export_experiment_workbook(config: dict, grid_output: dict) -> Path:
    """Export the G0-G4 tables, including attribution and multi-seed summaries."""
    tables_dir = Path(config["paths"]["tables_dir"])
    tables_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = tables_dir / "experiment_results.xlsx"
    used_sheets: set[str] = set()

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for sheet_name, key in (
            ("all_seed_runs", "runs"),
            ("seeds_summary", "seeds_summary"),
            ("rho_sweep", "rho_sweep"),
            ("ablation", "ablation"),
            ("significance", "significance"),
        ):
            frame = grid_output.get(key, pd.DataFrame())
            if not isinstance(frame, pd.DataFrame):
                frame = pd.DataFrame(frame)
            frame.to_excel(writer, sheet_name=_safe_sheet_name(sheet_name, used_sheets), index=False)

        attribution_by_scenario: dict[str, list[dict]] = {}
        exported_results: set[int] = set()
        for run_key, result in grid_output.get("detailed_runs", {}).items():
            attribution = result.get("attribution")
            if not attribution:
                continue
            # A cached result can be shared by several groups (e.g. S4 rho=1.0
            # m=3 in G2 and G4); export its per-round attribution only once.
            if id(result) in exported_results:
                continue
            exported_results.add(id(result))
            scenario = result["scenario"]
            for row in attribution.get("per_round", []):
                export_row = dict(row)
                export_row["flagged_clients"] = json.dumps(export_row.get("flagged_clients", []))
                export_row.update(
                    {
                        "run_key": run_key,
                        "seed": result["seed"],
                        "rho": result["rho"],
                        "malicious_clients": json.dumps(result["malicious_clients"]),
                        "precision_final": attribution.get("precision_cum"),
                        "recall_final": attribution.get("recall_cum"),
                        "time_to_detection": json.dumps(attribution.get("time_to_detection", {})),
                        "mean_time_to_detection": attribution.get("mean_time_to_detection"),
                        "false_positive_rate": attribution.get("false_positive_rate"),
                        "ecobee_false_positive": attribution.get("ecobee_false_positive"),
                        "audit_correctness": attribution.get("audit_correctness"),
                    }
                )
                attribution_by_scenario.setdefault(scenario, []).append(export_row)

        for scenario, rows in attribution_by_scenario.items():
            sheet_name = _safe_sheet_name(f"attribution_{scenario}", used_sheets)
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet_name, index=False)

    workbook = load_workbook(workbook_path)
    for sheet_name in workbook.sheetnames:
        _style_sheet(workbook[sheet_name])
    workbook.save(workbook_path)
    return workbook_path
